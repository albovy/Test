#!/usr/bin/env python3
"""Extrae la lista de equipos de un club de Clupik y la guarda en un Excel (.xlsx).

Fuente de datos: API pública de Clupik (no requiere iniciar sesión).

    GET https://api.clupik.com/clubs/<CLUB_ID>/teams

Club por defecto: Club Voleibol San Martiño (Dompavolei), ID 27574.

Uso:
    python extraer_equipos.py
    python extraer_equipos.py --club-id 27574 -o equipos.xlsx

Requisitos (ver requirements.txt):
    requests, openpyxl
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime, timezone

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

API_BASE = "https://api.clupik.com"
DEFAULT_CLUB_ID = "27574"  # Club Voleibol San Martiño (Dompavolei)
TIMEOUT = 30  # segundos

# Columnas del Excel: (encabezado, función que extrae el valor de cada equipo).
# Cada "registro" de la API es la pertenencia de un equipo al club; el detalle
# del equipo vive en la clave anidada "team".
COLUMNS: list[tuple[str, callable]] = [
    ("ID equipo",      lambda r, t, s: t.get("id")),
    ("Nombre",         lambda r, t, s: t.get("name")),
    ("Nombre visible", lambda r, t, s: t.get("editableName")),
    ("Descripción",    lambda r, t, s: t.get("description")),
    ("Temporada",      lambda r, t, s: s.get("title")),
    ("Sección",        lambda r, t, s: r.get("sectionId")),
    ("Visible",        lambda r, t, s: _si_no(t.get("visibility"))),
    ("Destacado",      lambda r, t, s: _si_no(t.get("isFeatured"))),
    ("Orden",          lambda r, t, s: t.get("order")),
    ("Nombre FEB",     lambda r, t, s: t.get("febName")),
    ("Escudo (URL)",   lambda r, t, s: t.get("shieldUrl")),
    ("ID club",        lambda r, t, s: r.get("clubId")),
    ("ID registro",    lambda r, t, s: r.get("id")),
]


def _si_no(value: object) -> str:
    """Convierte un booleano en texto legible en español."""
    return "Sí" if value else "No"


def fetch_teams(club_id: str) -> list[dict]:
    """Descarga la lista de equipos del club desde la API pública de Clupik."""
    url = f"{API_BASE}/clubs/{club_id}/teams"
    resp = requests.get(url, timeout=TIMEOUT, headers={"Accept": "application/json"})
    resp.raise_for_status()
    data = resp.json()
    if not isinstance(data, list):
        raise ValueError(f"Respuesta inesperada de la API (no es una lista): {type(data)}")
    return data


def build_rows(teams: list[dict]) -> list[list]:
    """Aplana cada equipo en una fila según COLUMNS."""
    rows = []
    for record in teams:
        team = record.get("team") or {}
        season = team.get("season") or {}
        rows.append([getter(record, team, season) for _, getter in COLUMNS])
    return rows


def write_xlsx(rows: list[list], club_id: str, output_path: str) -> None:
    """Escribe las filas en un libro de Excel con formato."""
    wb = Workbook()

    # --- Hoja principal: Equipos ---
    ws = wb.active
    ws.title = "Equipos"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    headers = [h for h, _ in COLUMNS]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append(row)

    # Inmoviliza encabezados y activa el filtro.
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(headers))
    ws.auto_filter.ref = f"A1:{last_col}{len(rows) + 1}"

    # Ajusta el ancho de cada columna al contenido (con un tope).
    for col_idx, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row in rows:
            value = row[col_idx - 1]
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    # --- Hoja de información / procedencia ---
    info = wb.create_sheet("Info")
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    for label, value in [
        ("Club ID", club_id),
        ("Fuente", f"{API_BASE}/clubs/{club_id}/teams"),
        ("Nº de equipos", len(rows)),
        ("Generado", generated),
    ]:
        info.append([label, value])
    info.column_dimensions["A"].width = 16
    info.column_dimensions["B"].width = 60
    for cell in info["A"]:
        cell.font = Font(bold=True)

    wb.save(output_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Extrae los equipos de un club de Clupik a Excel.")
    parser.add_argument("--club-id", default=DEFAULT_CLUB_ID,
                        help=f"ID del club en Clupik (por defecto: {DEFAULT_CLUB_ID}).")
    parser.add_argument("-o", "--output", default="equipos_dompavolei.xlsx",
                        help="Ruta del archivo Excel de salida.")
    args = parser.parse_args(argv)

    print(f"Descargando equipos del club {args.club_id}…")
    try:
        teams = fetch_teams(args.club_id)
    except requests.RequestException as exc:
        print(f"Error de red al consultar la API: {exc}", file=sys.stderr)
        return 1
    except ValueError as exc:
        print(f"Error en los datos recibidos: {exc}", file=sys.stderr)
        return 1

    rows = build_rows(teams)
    write_xlsx(rows, args.club_id, args.output)
    print(f"Listo: {len(rows)} equipos guardados en «{args.output}».")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
